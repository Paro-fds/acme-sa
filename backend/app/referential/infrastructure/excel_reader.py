"""US-501 CA-01 : lecture du classeur Excel du référentiel (onglets « Unités » et « Correspondances »)."""

import io
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.referential.domain.errors import UnreadableReferentialFile
from app.referential.domain.units import (
    MAPPINGS_SHEET,
    UNITS_SHEET,
    MappingRow,
    ReferentialFile,
    UnitRow,
    normalize,
)

UNIT_COLUMNS = {
    "code": "Code",
    "label": "Libellé officiel",
    "type": "Type",
    "parent": "Rattachée à",
    "since": "Rattachée depuis",
    "former": "Anciens libellés",
}
MAPPING_COLUMNS = {
    "column": "Colonne de l'export",
    "value": "Valeur dans l'export",
    "unit": "Code de l'unité",
}
REQUIRED_UNIT_COLUMNS = ("code", "label", "type", "parent")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()
    except ValueError as error:
        raise UnreadableReferentialFile(f"Date « {value} » illisible : écrivez-la JJ/MM/AAAA.") from error


def _rows(workbook, sheet_name: str, columns: dict[str, str], required: tuple[str, ...]):
    sheet = next((ws for ws in workbook.worksheets if normalize(ws.title) == normalize(sheet_name)), None)
    if sheet is None:
        raise UnreadableReferentialFile(f"Onglet « {sheet_name} » introuvable dans le classeur.")
    rows = sheet.iter_rows(values_only=True)
    header = [normalize(_text(cell)) for cell in next(rows, ())]
    index = {}
    for key, title in columns.items():
        if normalize(title) in header:
            index[key] = header.index(normalize(title))
        elif key in required:
            raise UnreadableReferentialFile(f"Colonne « {title} » introuvable dans l'onglet « {sheet_name} ».")
    for line, values in enumerate(rows, start=2):
        if not any(_text(value) for value in values):
            continue
        yield line, {key: values[position] if position < len(values) else None for key, position in index.items()}


class OpenpyxlReferentialReader:
    def read(self, content: bytes) -> ReferentialFile:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as error:
            raise UnreadableReferentialFile() from error
        try:
            units = tuple(
                UnitRow(
                    line=line,
                    code=_text(row["code"]).upper(),
                    label=_text(row["label"]),
                    type_label=_text(row["type"]),
                    parent_code=_text(row.get("parent")).upper(),
                    parent_since=_date(row.get("since")),
                    former_labels=tuple(
                        part.strip() for part in _text(row.get("former")).split(";") if part.strip()
                    ),
                )
                for line, row in _rows(workbook, UNITS_SHEET, UNIT_COLUMNS, REQUIRED_UNIT_COLUMNS)
            )
            mappings = tuple(
                MappingRow(
                    line=line,
                    column=_text(row["column"]),
                    value=_text(row["value"]),
                    unit_code=_text(row.get("unit")).upper(),
                )
                for line, row in _rows(workbook, MAPPINGS_SHEET, MAPPING_COLUMNS, ("column", "value"))
            )
        finally:
            workbook.close()
        return ReferentialFile(units, mappings)
