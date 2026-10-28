#!/usr/bin/env python3
"""Export mensuel Shivam vers le S3 de production lu par PAM.

En fin de mois, le partage reçoit deux rapports HTML enregistrés en ``.xls`` :
``Employee_PER_Accounts`` (balances PER) et ``Emp_Salary__Pay_HeadDetails``
(bulletins). Ce module les découpe en un objet JSON par employé.

Sans période, la lecture renvoie la dernière balance ou le dernier bulletin,
à la date du rapport. Avec un début et une fin (``AAAA-MM``), elle renvoie
le relevé de ces mois. Republier un même mois remplace cette photo : l'historique
des autres mois reste en place.

Les clés, sous le préfixe ``pam/balances`` par défaut :

- ``per/{CODE}.json``
- ``payslip/{CODE}.json``
- ``meta/imports.json`` (nombre d'employés et période, aucun nom ni montant)

PAM ne reçoit que ``s3:GetObject`` sur ``per/`` et ``payslip/``. L'export,
lui, écrit. Le code employé est la clé : un objet ne contient jamais un autre
employé.

Exemple, une fois ``PAM_S3_BUCKET`` renseigné dans l'environnement ou dans
``pam_export.env`` ::

    python export_shivam.py exporter --partage \\\\serveur\\shivam --s3 --env pam_export.env

``--sortie`` écrit le même arbre en local, sans AWS. ``--dry-run`` contrôle
les fichiers et n'écrit rien.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Protocol

CENT = Decimal("0.01")
PREFIXE_DEFAUT = "pam/balances"
FAMILLES = ("per", "payslip")
CHAMP_LISTE = {"per": "snapshots", "payslip": "bulletins"}
CHAMP_DATE = {"per": "date_rapport", "payslip": "date_paie"}
ROLES_PER = ("numero", "code", "nom", "ouverture", "balance", "interet")

# Le total Shivam est la somme des gains (crédit ISR compris) moins les
# retenues du salarié. Les parts employeur figurent sur le bulletin mais
# n'entrent pas dans ce total. Vérifié sur le relevé d'octobre 2026 :
# salaire et salaire contractuel ne sont jamais remplis ensemble.
NATURES = {
    "AA": "gain",
    "AB": "gain",
    "AC": "gain",
    "AE": "gain",
    "AF": "gain",
    "AG": "retenue",
    "AH": "retenue",
    "AI": "retenue",
    "AJ": "retenue",
    "AK": "employeur",
    "AL": "employeur",
    "AM": "employeur",
    "AN": "retenue",
    "AO": "retenue",
    "AP": "retenue",
    "AQ": "retenue",
    "AR": "gain",
    "AS": "gain",
    "AT": "retenue",
    "AU": "retenue",
    "AV": "employeur",
    "AW": "retenue",
    "AX": "gain",
    "AY": "retenue",
    "AZ": "gain",
    "BA": "retenue",
    "BB": "retenue",
}

MOTIF_RAPPORT = re.compile(r"Employee_PER_Accounts|Emp_Salary__Pay_HeadDetails", re.IGNORECASE)
MOTIF_RUBRIQUE = re.compile(r"([A-Za-z]{2})\s*\((.*)\)\s*\Z")
SUFFIXES = {".xls", ".xlsx", ".html", ".htm"}


class ErreurExport(Exception):
    """Le rapport ou la commande ne peuvent pas être publiés tels quels."""


class Magasin(Protocol):
    def lire(self, famille: str, code: str) -> dict | None: ...

    def ecrire(self, famille: str, code: str, document: dict) -> None: ...

    def effacer(self, famille: str, code: str) -> None: ...

    def codes(self, famille: str) -> list[str]: ...

    def lire_meta(self) -> dict | None: ...

    def ecrire_meta(self, document: dict) -> None: ...


@dataclass(frozen=True)
class Morceau:
    code: str
    periode: str
    date: str
    source: str
    contenu: dict


@dataclass(frozen=True)
class Analyse:
    famille: str
    source: str
    morceaux: tuple[Morceau, ...]


@dataclass(frozen=True)
class LigneBilan:
    famille: str
    periode: str
    date: str
    employes: int


@dataclass(frozen=True)
class Bilan:
    lignes: tuple[LigneBilan, ...]

    def texte(self) -> str:
        if not self.lignes:
            return "Aucun rapport."
        morceaux = []
        for ligne in self.lignes:
            libelle = "PER" if ligne.famille == "per" else "Bulletins"
            nombre = "1 employé" if ligne.employes == 1 else f"{ligne.employes} employés"
            morceaux.append(f"{libelle} {ligne.periode} (au {date_fr(ligne.date)}) : {nombre}.")
        return " ".join(morceaux)


def sans_accent(valeur: str) -> str:
    decompose = unicodedata.normalize("NFD", valeur)
    return "".join(caractere for caractere in decompose if unicodedata.category(caractere) != "Mn")


def montant_texte(valeur: Decimal) -> str:
    return f"{valeur.quantize(CENT, rounding=ROUND_HALF_UP):.2f}"


def parse_montant(valeur: object) -> Decimal:
    """Lit un montant Shivam : ``1.234,50``, ``,00`` ou ``000``."""
    if isinstance(valeur, bool) or valeur is None:
        return Decimal("0.00")
    if isinstance(valeur, Decimal):
        return valeur.quantize(CENT, rounding=ROUND_HALF_UP)
    if isinstance(valeur, int):
        return Decimal(valeur).quantize(CENT, rounding=ROUND_HALF_UP)
    if isinstance(valeur, float):
        return Decimal(f"{valeur:.2f}").quantize(CENT, rounding=ROUND_HALF_UP)
    brut = str(valeur).replace("\u00a0", "").replace(" ", "").strip()
    if brut.lower() in {"", "nan", "none", "-", "–", "—"}:
        return Decimal("0.00")
    negatif = brut.startswith("-") or (brut.startswith("(") and brut.endswith(")"))
    brut = brut.strip("()").lstrip("+").lstrip("-")
    if "," in brut and "." in brut:
        if brut.rfind(",") > brut.rfind("."):
            brut = brut.replace(".", "").replace(",", ".")
        else:
            brut = brut.replace(",", "")
    elif "," in brut:
        brut = brut.replace(".", "").replace(",", ".")
    elif brut.count(".") > 1 or re.fullmatch(r"\d{1,3}(\.\d{3})+", brut):
        brut = brut.replace(".", "")
    try:
        nombre = Decimal(brut)
    except InvalidOperation as exc:
        raise ErreurExport(f"Montant illisible : {str(valeur)[:40]!r}.") from exc
    if negatif:
        nombre = -nombre
    return nombre.quantize(CENT, rounding=ROUND_HALF_UP)


def parse_date(valeur: object) -> str:
    texte = " ".join(str(valeur or "").split())
    if not texte:
        raise ErreurExport("Date manquante dans le rapport.")
    jour = texte.split()[0]
    correspondance = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", jour)
    if correspondance:
        jour_n, mois, annee = (int(partie) for partie in correspondance.groups())
    else:
        iso = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", jour)
        if not iso:
            raise ErreurExport(f"Date illisible : {texte[:40]!r}.")
        annee, mois, jour_n = (int(partie) for partie in iso.groups())
    try:
        return date(annee, mois, jour_n).isoformat()
    except ValueError as exc:
        raise ErreurExport(f"Date illisible : {texte[:40]!r}.") from exc


def date_fr(iso: str) -> str:
    annee, mois, jour = iso.split("-")
    return f"{jour}/{mois}/{annee}"


def normaliser_code(valeur: object) -> str:
    """Code comparable : majuscules, sans espace ni apostrophe.

    Shivam a au moins un code saisi avec un accent grave (``D`JO0818``).
    L'apostrophe n'entre pas dans la clé S3 ; elle est retirée à la lecture
    comme à l'écriture, donc les deux saisies désignent le même objet.
    """
    code = "".join(str(valeur or "").split()).upper()
    for apostrophe in ("`", "'", "’", "´", "‘"):
        code = code.replace(apostrophe, "")
    if (
        code == "TOTAL"
        or not re.fullmatch(r"[A-Z0-9]{4,16}", code)
        or not any(caractere.isalpha() for caractere in code)
    ):
        raise ErreurExport(f"Code employé illisible : {code[:32]!r}.")
    return code


def valider_periode(valeur: str | None) -> str | None:
    if valeur is None or valeur == "":
        return None
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", valeur):
        raise ErreurExport("Période attendue au format AAAA-MM.")
    return valeur


def est_total(valeur: object) -> bool:
    return "".join(str(valeur or "").split()).upper() == "TOTAL"


def est_ligne_synthese(valeur: object) -> bool:
    """Ligne de total Shivam, ou un montant tombé dans la colonne du code."""
    plat = sans_accent(" ".join(str(valeur or "").split())).upper()
    if not plat or plat in {"TOTAL", "TOTAUX"} or plat.startswith("TOTAL"):
        return True
    return not any(caractere.isalpha() for caractere in plat)


def decoder(raw: bytes) -> str:
    for encodage in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encodage)
        except UnicodeDecodeError:
            continue
    return raw.decode("cp1252", errors="replace")


def date_depuis_nom(chemin: Path) -> str | None:
    correspondance = re.search(r"(20\d{2})(\d{2})(\d{2})", chemin.name)
    if not correspondance:
        return None
    annee, mois, jour = (int(partie) for partie in correspondance.groups())
    try:
        return date(annee, mois, jour).isoformat()
    except ValueError:
        return None


def date_imprime(texte: str, chemin: Path) -> str:
    plat = sans_accent(texte)[:8000]
    correspondance = re.search(
        r"imprime\s+le\s*-?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{4})",
        plat,
        re.IGNORECASE,
    )
    if correspondance:
        return parse_date(correspondance.group(1))
    trouvee = date_depuis_nom(chemin)
    if trouvee:
        return trouvee
    raise ErreurExport(f"Date du rapport introuvable : {chemin.name}")


def lire_tables(chemin: Path) -> tuple[str, list[list[list[str]]]]:
    if not chemin.is_file():
        raise ErreurExport(f"Fichier introuvable : {chemin.name}")
    raw = chemin.read_bytes()
    debut = raw.lstrip()[:800].lower()
    if b"<html" in debut or b"<table" in debut or debut.startswith(b"<"):
        texte = decoder(raw)
        return texte, _tables_html(texte)
    if raw[:2] == b"PK":
        return "", _tables_xlsx(chemin)
    if raw.startswith(b"\xd0\xcf\x11\xe0"):
        return "", _tables_xls(chemin)
    raise ErreurExport(
        f"{chemin.name} n'est ni le HTML Shivam ni un classeur. "
        "Enregistrez le rapport tel qu'il sort de Shivam."
    )


def _tables_html(texte: str) -> list[list[list[str]]]:
    from lxml import html

    document = html.fromstring(texte)
    tables: list[list[list[str]]] = []
    for table in document.xpath("//table"):
        lignes = []
        for ligne in table.xpath(".//tr"):
            cellules = ligne.xpath("./th|./td")
            if not cellules:
                continue
            lignes.append([" ".join("".join(cellule.itertext()).split()) for cellule in cellules])
        if lignes:
            tables.append(lignes)
    if not tables:
        raise ErreurExport("Aucune table dans le rapport HTML.")
    return tables


def _texte_cellule(valeur: object) -> str:
    if valeur is None:
        return ""
    return " ".join(str(valeur).split())


def _tables_xlsx(chemin: Path) -> list[list[list[str]]]:
    from openpyxl import load_workbook

    classeur = load_workbook(chemin, read_only=True, data_only=True)
    tables: list[list[list[str]]] = []
    try:
        for feuille in classeur.worksheets:
            lignes = [
                [_texte_cellule(cellule) for cellule in ligne]
                for ligne in feuille.iter_rows(values_only=True)
                if any(cellule is not None and str(cellule).strip() for cellule in ligne)
            ]
            if lignes:
                tables.append(lignes)
    finally:
        classeur.close()
    return tables


def _tables_xls(chemin: Path) -> list[list[list[str]]]:
    import xlrd

    classeur = xlrd.open_workbook(chemin)
    tables: list[list[list[str]]] = []
    for feuille in classeur.sheets():
        lignes = []
        for index in range(feuille.nrows):
            valeurs = [_texte_cellule(feuille.cell_value(index, colonne)) for colonne in range(feuille.ncols)]
            if any(valeurs):
                lignes.append(valeurs)
        if lignes:
            tables.append(lignes)
    return tables


def _role_per(cellule: str) -> str | None:
    nom = " ".join(sans_accent(cellule).lower().replace(".", " ").split())
    if "compte" in nom and "per" in nom:
        return "numero"
    if "code" in nom and "empl" in nom:
        return "code"
    if nom.startswith("nom"):
        return "nom"
    # Shivam écrit « Date Overture », sans le u de « ouverture ».
    if "ouvert" in nom or "overt" in nom:
        return "ouverture"
    if nom == "balance" or nom.startswith("balance"):
        return "balance"
    if "inter" in nom:
        return "interet"
    return None


def _indexer_per(entete: list[str]) -> dict[str, int] | None:
    roles: dict[str, int] = {}
    inconnues: list[str] = []
    for index, cellule in enumerate(entete):
        if not cellule:
            continue
        role = _role_per(cellule)
        if role is None:
            inconnues.append(cellule)
        elif role in roles:
            raise ErreurExport(f"Colonne PER en double : {cellule}.")
        else:
            roles[role] = index
    if "numero" not in roles and "code" not in roles:
        return None
    manquantes = [role for role in ROLES_PER if role not in roles]
    if manquantes or inconnues:
        raise ErreurExport(
            "En-tête PER inattendu. Colonnes manquantes : "
            f"{', '.join(manquantes) or 'aucune'} ; inconnues : {', '.join(inconnues) or 'aucune'}."
        )
    return roles


def _dedoublonner_comptes(code: str, comptes: list[dict]) -> list[dict]:
    par_numero: dict[str, dict] = {}
    for compte in comptes:
        numero = compte["numero"]
        if not numero:
            raise ErreurExport(f"Numéro de compte PER manquant pour {code}.")
        ancien = par_numero.get(numero)
        if ancien is None:
            par_numero[numero] = compte
        elif ancien != compte:
            raise ErreurExport(
                f"Le compte {numero} du code {code} apparaît deux fois avec des valeurs différentes."
            )
    return [par_numero[numero] for numero in sorted(par_numero)]


def parser_per(tables: list[list[list[str]]], date_rapport: str, source: str) -> tuple[Morceau, ...]:
    comptes: dict[str, list[dict]] = {}
    noms: dict[str, str] = {}
    for table in tables:
        if not table:
            continue
        colonnes = _indexer_per(table[0])
        if colonnes is None:
            continue
        largeur = max(colonnes.values()) + 1
        for brute in table[1:]:
            ligne = brute + [""] * (largeur - len(brute))
            if not any(ligne) or est_ligne_synthese(ligne[colonnes["code"]]):
                continue
            code = normaliser_code(ligne[colonnes["code"]])
            nom = " ".join(ligne[colonnes["nom"]].split())
            if nom and code in noms and noms[code] != nom:
                raise ErreurExport(f"Le code {code} a plusieurs noms dans le rapport PER.")
            if nom:
                noms[code] = nom
            ouverture = ligne[colonnes["ouverture"]]
            comptes.setdefault(code, []).append(
                {
                    "numero": ligne[colonnes["numero"]],
                    "date_ouverture": parse_date(ouverture) if ouverture else "",
                    "balance": montant_texte(parse_montant(ligne[colonnes["balance"]])),
                    "interet_cumule": montant_texte(parse_montant(ligne[colonnes["interet"]])),
                }
            )
    if not comptes:
        raise ErreurExport(f"Aucun compte PER dans {source}.")
    periode = date_rapport[:7]
    morceaux: list[Morceau] = []
    for code in sorted(comptes):
        liste = _dedoublonner_comptes(code, comptes[code])
        balance = sum((Decimal(compte["balance"]) for compte in liste), Decimal("0"))
        interet = sum((Decimal(compte["interet_cumule"]) for compte in liste), Decimal("0"))
        contenu = {
            "periode": periode,
            "date_rapport": date_rapport,
            "nom": noms.get(code, ""),
            "balance": montant_texte(balance),
            "interet_cumule": montant_texte(interet),
            "comptes": liste,
            "source": source,
        }
        morceaux.append(Morceau(code, periode, date_rapport, source, contenu))
    return tuple(morceaux)


def _role_paie(cellule: str) -> str | None:
    nom = " ".join(sans_accent(cellule).lower().split())
    if nom == "date":
        return "date"
    if nom in {"employee id", "employeeid", "code employe"}:
        return "code"
    if nom in {"employee name", "nom employee", "nom employe"}:
        return "nom"
    if nom == "total":
        return "total"
    return None


def _indexer_paie(entete: list[str]) -> tuple[dict[str, int], list[tuple[int, str, str, str]]] | None:
    fixes: dict[str, int] = {}
    rubriques: list[tuple[int, str, str, str]] = []
    inconnues: list[str] = []
    for index, cellule in enumerate(entete):
        if not str(cellule).strip():
            continue
        role = _role_paie(cellule)
        if role:
            if role in fixes:
                raise ErreurExport(f"Colonne de paie en double : {cellule}.")
            fixes[role] = index
            continue
        correspondance = MOTIF_RUBRIQUE.fullmatch(cellule.strip())
        if correspondance:
            code = correspondance.group(1).upper()
            libelle = " ".join(correspondance.group(2).split())
            nature = NATURES.get(code)
            if nature is None:
                inconnues.append(code)
            else:
                rubriques.append((index, code, libelle, nature))
            continue
        inconnues.append(cellule)
    if "code" not in fixes and "date" not in fixes:
        return None
    manquantes = [role for role in ("date", "code", "nom", "total") if role not in fixes]
    if manquantes or inconnues or not rubriques:
        raise ErreurExport(
            "En-tête de bulletin inattendu. Manquantes : "
            f"{', '.join(manquantes) or 'aucune'} ; inconnues : {', '.join(inconnues) or 'aucune'}."
        )
    return fixes, rubriques


def parser_paie(tables: list[list[list[str]]], source: str) -> tuple[Morceau, ...]:
    index = None
    lignes: list[list[str]] = []
    for table in tables:
        if not table:
            continue
        trouve = _indexer_paie(table[0])
        if trouve is None:
            continue
        if index is not None:
            raise ErreurExport(f"Plusieurs tableaux de paie dans {source}.")
        index = trouve
        lignes = table[1:]
    if index is None:
        raise ErreurExport(f"Aucun bulletin dans {source}.")
    fixes, rubriques = index
    vus: set[tuple[str, str]] = set()
    morceaux: list[Morceau] = []
    largeur = max([*fixes.values(), *[rubrique[0] for rubrique in rubriques]]) + 1
    for brute in lignes:
        ligne = brute + [""] * (largeur - len(brute))
        if (
            not any(ligne)
            or est_total(ligne[fixes["date"]])
            or est_ligne_synthese(ligne[fixes["code"]])
        ):
            continue
        code = normaliser_code(ligne[fixes["code"]])
        date_paie = parse_date(ligne[fixes["date"]])
        periode = date_paie[:7]
        if (code, periode) in vus:
            raise ErreurExport(f"Code employé en double dans le bulletin : {code}.")
        vus.add((code, periode))
        detail = []
        gains = Decimal("0")
        retenues = Decimal("0")
        employeur = Decimal("0")
        for position, code_rubrique, libelle, nature in rubriques:
            montant = parse_montant(ligne[position] if position < len(ligne) else "")
            detail.append(
                {
                    "code": code_rubrique,
                    "libelle": libelle,
                    "nature": nature,
                    "montant": montant_texte(montant),
                }
            )
            if nature == "gain":
                gains += montant
            elif nature == "retenue":
                retenues += montant
            else:
                employeur += montant
        total = parse_montant(ligne[fixes["total"]])
        if gains - retenues != total:
            raise ErreurExport(f"Le total du bulletin ne correspond pas aux rubriques ({code}).")
        contenu = {
            "periode": periode,
            "date_paie": date_paie,
            "nom": " ".join(ligne[fixes["nom"]].split()),
            "lignes": detail,
            "gains": montant_texte(gains),
            "retenues": montant_texte(retenues),
            "charges_employeur": montant_texte(employeur),
            "total": montant_texte(total),
            "source": source,
        }
        morceaux.append(Morceau(code, periode, date_paie, source, contenu))
    if not morceaux:
        raise ErreurExport(f"Aucun bulletin dans {source}.")
    return tuple(morceaux)


def analyser(chemin: Path) -> Analyse:
    """Lit un rapport Shivam et le découpe, sans rien écrire."""
    texte, tables = lire_tables(chemin)
    source = chemin.name
    if any(_indexer_per(table[0]) is not None for table in tables if table):
        return Analyse("per", source, parser_per(tables, date_imprime(texte, chemin), source))
    if any(_indexer_paie(table[0]) is not None for table in tables if table):
        return Analyse("payslip", source, parser_paie(tables, source))
    raise ErreurExport(f"Rapport ni PER ni bulletin : {source}.")


def regrouper(analyses: list[Analyse]) -> dict[str, dict[str, dict[str, Morceau]]]:
    """famille → période → code → morceau le plus récent."""
    sortie: dict[str, dict[str, dict[str, Morceau]]] = {famille: {} for famille in FAMILLES}
    for analyse in analyses:
        for morceau in analyse.morceaux:
            seau = sortie[analyse.famille].setdefault(morceau.periode, {})
            ancien = seau.get(morceau.code)
            if ancien is None or morceau.date >= ancien.date:
                seau[morceau.code] = morceau
    return sortie


def fusionner(existant: dict | None, code: str, contenu: dict, famille: str) -> dict:
    champ = CHAMP_LISTE[famille]
    champ_date = CHAMP_DATE[famille]
    if existant is not None and existant.get("schema") != 1:
        raise ErreurExport(f"Schéma inattendu pour l'objet {famille} de {code}.")
    par_periode = {item["periode"]: item for item in (existant or {}).get(champ, [])}
    ancien = par_periode.get(contenu["periode"])
    if ancien is None or contenu[champ_date] >= ancien[champ_date]:
        par_periode[contenu["periode"]] = contenu
    items = sorted(par_periode.values(), key=lambda item: (item["periode"], item[champ_date]))
    nom = max(items, key=lambda item: (item["periode"], item[champ_date]))["nom"]
    return {"schema": 1, "code_employe": code, "nom": nom, champ: items}


def retirer_periode(document: dict, famille: str, periode: str) -> dict | None:
    champ = CHAMP_LISTE[famille]
    champ_date = CHAMP_DATE[famille]
    reste = [item for item in document.get(champ, []) if item["periode"] != periode]
    if not reste:
        return None
    nom = max(reste, key=lambda item: (item["periode"], item[champ_date]))["nom"]
    return {"schema": 1, "code_employe": document["code_employe"], "nom": nom, champ: reste}


def population(magasin: Magasin, famille: str, periode: str) -> int:
    champ = CHAMP_LISTE[famille]
    nombre = 0
    for code in magasin.codes(famille):
        document = magasin.lire(famille, code)
        if document and any(item["periode"] == periode for item in document.get(champ, [])):
            nombre += 1
    return nombre


def _controler(magasin: Magasin, groupes: dict, complet: bool, forcer: bool) -> None:
    if not complet:
        return
    for famille, par_periode in groupes.items():
        for periode, par_code in par_periode.items():
            deja = population(magasin, famille, periode)
            if deja and len(par_code) * 2 < deja and not forcer:
                raise ErreurExport(
                    f"Le rapport {famille} {periode} contient {len(par_code)} employés, "
                    f"contre {deja} déjà publiés. Relancez avec --forcer si le fichier est bien complet."
                )


def _meta(magasin: Magasin, groupes: dict, complet: bool) -> None:
    meta = magasin.lire_meta() or {"schema": 1, "imports": []}
    if meta.get("schema") != 1:
        raise ErreurExport("Schéma inattendu pour meta/imports.json.")
    imports = list(meta.get("imports") or [])
    maintenant = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    for famille, par_periode in groupes.items():
        for periode, par_code in par_periode.items():
            dates = sorted({morceau.date for morceau in par_code.values()})
            sources = sorted({morceau.source for morceau in par_code.values()})
            imports = [
                entree
                for entree in imports
                if not (entree.get("famille") == famille and entree.get("periode") == periode)
            ]
            imports.append(
                {
                    "famille": famille,
                    "periode": periode,
                    "date": dates[-1],
                    "employes": len(par_code),
                    "complet": complet,
                    "sources": sources,
                    "importe_le": maintenant,
                }
            )
    imports.sort(key=lambda entree: (entree["famille"], entree["periode"]))
    magasin.ecrire_meta({"schema": 1, "imports": imports})


def exporter(
    chemins: list[Path],
    magasin: Magasin | None = None,
    *,
    complet: bool = True,
    forcer: bool = False,
) -> Bilan:
    """Publie les rapports. Sans magasin, contrôle seulement (dry-run)."""
    if not chemins:
        raise ErreurExport("Aucun rapport Shivam trouvé.")
    analyses = [analyser(chemin) for chemin in chemins]
    groupes = regrouper(analyses)
    lignes: list[LigneBilan] = []
    for famille in FAMILLES:
        for periode, par_code in sorted(groupes[famille].items()):
            date_periode = max(morceau.date for morceau in par_code.values())
            lignes.append(LigneBilan(famille, periode, date_periode, len(par_code)))
    bilan = Bilan(tuple(lignes))
    if magasin is None:
        return bilan
    _controler(magasin, groupes, complet, forcer)
    for famille in FAMILLES:
        for periode, par_code in sorted(groupes[famille].items()):
            for code, morceau in sorted(par_code.items()):
                document = fusionner(magasin.lire(famille, code), code, morceau.contenu, famille)
                magasin.ecrire(famille, code, document)
            if not complet:
                continue
            for code in list(magasin.codes(famille)):
                if code in par_code:
                    continue
                actuel = magasin.lire(famille, code)
                if actuel is None or not any(
                    item["periode"] == periode for item in actuel.get(CHAMP_LISTE[famille], [])
                ):
                    continue
                reste = retirer_periode(actuel, famille, periode)
                if reste is None:
                    magasin.effacer(famille, code)
                else:
                    magasin.ecrire(famille, code, reste)
    _meta(magasin, groupes, complet)
    return bilan


def consulter(
    magasin: Magasin,
    famille: str,
    code: str,
    debut: str | None = None,
    fin: str | None = None,
) -> dict | None:
    """Dernière photo, ou relevé si ``debut`` ou ``fin`` est fourni.

    ``None`` signifie que ce code n'a pas d'objet. Un relevé vide signifie
    que l'employé existe, mais pas sur les mois demandés.
    """
    if famille not in FAMILLES:
        raise ErreurExport("Famille attendue : per ou payslip.")
    debut = valider_periode(debut)
    fin = valider_periode(fin)
    if debut and fin and debut > fin:
        raise ErreurExport("Le début de période est postérieur à la fin.")
    code_normalise = normaliser_code(code)
    document = magasin.lire(famille, code_normalise)
    if not document:
        return None
    champ = CHAMP_LISTE[famille]
    champ_date = CHAMP_DATE[famille]
    items = list(document.get(champ) or [])
    if not items:
        return None
    if debut is None and fin is None:
        choisi = max(items, key=lambda item: (item["periode"], item[champ_date]))
        return _vue(famille, code_normalise, choisi, "courante")
    choisis = [
        item
        for item in items
        if (debut is None or item["periode"] >= debut) and (fin is None or item["periode"] <= fin)
    ]
    choisis.sort(key=lambda item: (item["periode"], item[champ_date]))
    return {
        "famille": famille,
        "mode": "releve",
        "code_employe": code_normalise,
        "nom": document.get("nom", ""),
        "debut": debut,
        "fin": fin,
        "periodes": [_periode(item) for item in choisis],
    }


def _periode(item: dict) -> dict:
    return {cle: valeur for cle, valeur in item.items() if cle != "source"}


def _vue(famille: str, code: str, item: dict, mode: str) -> dict:
    vue = _periode(item)
    vue["famille"] = famille
    vue["mode"] = mode
    vue["code_employe"] = code
    return vue


def _json(document: dict) -> bytes:
    return (json.dumps(document, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def _lire_json(payload: bytes) -> dict:
    document = json.loads(payload.decode("utf-8"))
    if not isinstance(document, dict):
        raise ErreurExport("Objet JSON inattendu.")
    return document


class MagasinLocal:
    """Arbre de fichiers calqué sur les clés S3. Sert aux essais et au secours."""

    def __init__(self, racine: Path):
        self.racine = Path(racine)

    def _chemin(self, *parties: str) -> Path:
        return self.racine.joinpath(*parties)

    def lire(self, famille: str, code: str) -> dict | None:
        chemin = self._chemin(famille, f"{code}.json")
        if not chemin.is_file():
            return None
        return _lire_json(chemin.read_bytes())

    def ecrire(self, famille: str, code: str, document: dict) -> None:
        chemin = self._chemin(famille, f"{code}.json")
        chemin.parent.mkdir(parents=True, exist_ok=True)
        temporaire = chemin.with_suffix(".json.tmp")
        temporaire.write_bytes(_json(document))
        temporaire.replace(chemin)

    def effacer(self, famille: str, code: str) -> None:
        chemin = self._chemin(famille, f"{code}.json")
        if chemin.is_file():
            chemin.unlink()

    def codes(self, famille: str) -> list[str]:
        dossier = self._chemin(famille)
        if not dossier.is_dir():
            return []
        return sorted(chemin.stem for chemin in dossier.glob("*.json"))

    def lire_meta(self) -> dict | None:
        chemin = self._chemin("meta", "imports.json")
        if not chemin.is_file():
            return None
        return _lire_json(chemin.read_bytes())

    def ecrire_meta(self, document: dict) -> None:
        self.ecrire("meta", "imports", document)


class _Introuvable(Exception):
    def __init__(self, cle: str):
        super().__init__(cle)
        self.response = {"Error": {"Code": "NoSuchKey"}}


class MagasinS3:
    """Objets JSON du bucket de production. Le client peut être injecté pour les tests."""

    def __init__(
        self,
        bucket: str,
        prefixe: str = PREFIXE_DEFAUT,
        *,
        region: str | None = None,
        sse: str = "AES256",
        client: object | None = None,
    ):
        self.bucket = bucket.strip()
        self.prefixe = prefixe.strip("/")
        self.region = region
        self.sse = sse
        self._client = client
        if not self.bucket:
            raise ErreurExport("PAM_S3_BUCKET est vide.")

    def client(self):
        if self._client is None:
            import boto3

            arguments = {"region_name": self.region} if self.region else {}
            self._client = boto3.client("s3", **arguments)
        return self._client

    def _cle(self, *parties: str) -> str:
        morceaux = [self.prefixe] if self.prefixe else []
        morceaux.extend(parties)
        return "/".join(morceaux)

    def _absent(self, exc: Exception) -> bool:
        code = getattr(exc, "response", {}).get("Error", {}).get("Code", "")
        return code in {"404", "NoSuchKey", "NotFound"}

    def lire(self, famille: str, code: str) -> dict | None:
        try:
            reponse = self.client().get_object(Bucket=self.bucket, Key=self._cle(famille, f"{code}.json"))
        except Exception as exc:
            if self._absent(exc):
                return None
            raise
        corps = reponse["Body"].read()
        if hasattr(reponse["Body"], "close"):
            reponse["Body"].close()
        return _lire_json(corps)

    def ecrire(self, famille: str, code: str, document: dict) -> None:
        arguments = {
            "Bucket": self.bucket,
            "Key": self._cle(famille, f"{code}.json"),
            "Body": _json(document),
            "ContentType": "application/json; charset=utf-8",
        }
        if self.sse and self.sse.lower() != "none":
            arguments["ServerSideEncryption"] = self.sse
        self.client().put_object(**arguments)

    def effacer(self, famille: str, code: str) -> None:
        self.client().delete_object(Bucket=self.bucket, Key=self._cle(famille, f"{code}.json"))

    def codes(self, famille: str) -> list[str]:
        prefixe = self._cle(famille, "")
        jeton = None
        codes: list[str] = []
        while True:
            arguments: dict = {"Bucket": self.bucket, "Prefix": prefixe}
            if jeton:
                arguments["ContinuationToken"] = jeton
            reponse = self.client().list_objects_v2(**arguments)
            for objet in reponse.get("Contents") or []:
                cle = objet["Key"]
                nom = cle.rsplit("/", 1)[-1]
                if nom.endswith(".json"):
                    codes.append(nom[: -len(".json")])
            if not reponse.get("IsTruncated"):
                break
            jeton = reponse.get("NextContinuationToken")
        return sorted(set(codes))

    def lire_meta(self) -> dict | None:
        return self.lire("meta", "imports")

    def ecrire_meta(self, document: dict) -> None:
        self.ecrire("meta", "imports", document)


class MagasinDouble:
    """Écrit dans les deux magasins. La lecture suit le premier (S3 s'il est là)."""

    def __init__(self, principal: Magasin, copie: Magasin):
        self.principal = principal
        self.copie = copie

    def lire(self, famille: str, code: str) -> dict | None:
        return self.principal.lire(famille, code)

    def ecrire(self, famille: str, code: str, document: dict) -> None:
        self.principal.ecrire(famille, code, document)
        self.copie.ecrire(famille, code, document)

    def effacer(self, famille: str, code: str) -> None:
        self.principal.effacer(famille, code)
        self.copie.effacer(famille, code)

    def codes(self, famille: str) -> list[str]:
        return self.principal.codes(famille)

    def lire_meta(self) -> dict | None:
        return self.principal.lire_meta()

    def ecrire_meta(self, document: dict) -> None:
        self.principal.ecrire_meta(document)
        self.copie.ecrire_meta(document)


def trouver_rapports(dossier: Path) -> list[Path]:
    if not dossier.is_dir():
        raise ErreurExport(f"Dossier introuvable : {dossier}")
    trouves = [
        chemin
        for chemin in dossier.rglob("*")
        if chemin.is_file()
        and not chemin.name.startswith("~$")
        and chemin.suffix.lower() in SUFFIXES
        and MOTIF_RAPPORT.search(chemin.name)
    ]
    return sorted(trouves, key=lambda chemin: (chemin.stat().st_mtime, chemin.name))


def collecter(partage: Path | None, fichiers: list[Path]) -> list[Path]:
    chemins: list[Path] = []
    if partage is not None:
        chemins.extend(trouver_rapports(partage))
    for fichier in fichiers:
        if not fichier.is_file():
            raise ErreurExport(f"Fichier introuvable : {fichier}")
        chemins.append(fichier)
    uniques = {chemin.resolve(): chemin for chemin in chemins}
    return sorted(uniques.values(), key=lambda chemin: (chemin.stat().st_mtime, chemin.name))


def charger_env(chemin: Path) -> None:
    if not chemin.is_file():
        raise ErreurExport(f"Fichier d'environnement introuvable : {chemin}")
    for ligne in chemin.read_text(encoding="utf-8-sig").splitlines():
        ligne = ligne.strip()
        if not ligne or ligne.startswith("#") or "=" not in ligne:
            continue
        cle, valeur = ligne.split("=", 1)
        os.environ.setdefault(cle.strip(), valeur.strip().strip('"').strip("'"))


def _magasin_s3() -> MagasinS3:
    bucket = os.environ.get("PAM_S3_BUCKET", "").strip()
    if not bucket:
        raise ErreurExport(
            "PAM_S3_BUCKET est vide. Indiquez le bucket de production des balances avant d'exporter."
        )
    prefixe = os.environ.get("PAM_S3_PREFIX", PREFIXE_DEFAUT).strip() or PREFIXE_DEFAUT
    region = os.environ.get("PAM_S3_REGION", "us-east-1").strip() or None
    sse = os.environ.get("PAM_S3_SSE", "AES256").strip() or "AES256"
    return MagasinS3(bucket, prefixe, region=region, sse=sse)


def magasin_de_commande(sortie: Path | None, vers_s3: bool) -> Magasin:
    if vers_s3 and sortie is not None:
        return MagasinDouble(_magasin_s3(), MagasinLocal(sortie))
    if vers_s3:
        return _magasin_s3()
    if sortie is None:
        raise ErreurExport("Indiquez --sortie, --s3, ou --dry-run.")
    return MagasinLocal(sortie)


def _stdio_utf8() -> None:
    for flux in (sys.stdout, sys.stderr):
        reconfigurer = getattr(flux, "reconfigure", None)
        if reconfigurer is None:
            continue
        try:
            reconfigurer(encoding="utf-8")
        except (OSError, ValueError):
            continue


def _preparer(args: argparse.Namespace) -> None:
    if args.env:
        charger_env(Path(args.env))
    if args.commande != "exporter":
        return
    if args.partage is None and not args.fichier:
        partage = os.environ.get("PAM_SHARE_DIR", "").strip()
        if partage:
            args.partage = partage


def _commande_exporter(args: argparse.Namespace) -> int:
    chemins = collecter(Path(args.partage) if args.partage else None, [Path(fichier) for fichier in args.fichier])
    magasin = None if args.dry_run else magasin_de_commande(
        Path(args.sortie) if args.sortie else None,
        args.s3,
    )
    bilan = exporter(chemins, magasin, complet=args.complet, forcer=args.forcer)
    suite = "Rien n'a été écrit." if args.dry_run else "Publication terminée."
    print(f"{bilan.texte()} {suite}")
    return 0


def _commande_consulter(args: argparse.Namespace) -> int:
    magasin = magasin_de_commande(Path(args.sortie) if args.sortie else None, args.s3)
    vue = consulter(magasin, args.famille, args.code, args.debut, args.fin)
    if vue is None:
        print("Aucune donnée pour ce code dans les rapports importés.", file=sys.stderr)
        return 2
    json.dump(vue, sys.stdout, ensure_ascii=False, indent=2)
    print()
    return 0


def construire_analyseur() -> argparse.ArgumentParser:
    analyseur = argparse.ArgumentParser(
        description="Découpe les rapports Shivam PER et bulletins en un objet par employé."
    )
    analyseur.add_argument("--env", help="Fichier KEY=VALEUR. L'environnement du processus l'emporte.")
    sous = analyseur.add_subparsers(dest="commande", required=True)

    exporter_cmd = sous.add_parser("exporter", help="Lire le partage (ou des fichiers) et publier.")
    exporter_cmd.add_argument("--partage", help="Dossier du partage où Shivam dépose les rapports.")
    exporter_cmd.add_argument("--fichier", action="append", default=[], help="Rapport explicite. Répétable.")
    exporter_cmd.add_argument("--sortie", help="Dossier local calqué sur les clés S3.")
    exporter_cmd.add_argument("--s3", action="store_true", help="Écrire dans le bucket PAM_S3_BUCKET.")
    exporter_cmd.add_argument("--dry-run", action="store_true", help="Contrôler les fichiers sans écrire.")
    exporter_cmd.add_argument(
        "--complet",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Cette période remplace la population déjà publiée (défaut). --no-complet ajoute sans retirer.",
    )
    exporter_cmd.add_argument(
        "--forcer",
        action="store_true",
        help="Publier même si le rapport a moins de la moitié des employés déjà présents.",
    )

    consulter_cmd = sous.add_parser("consulter", help="Lire la dernière photo ou un relevé.")
    consulter_cmd.add_argument("--code", required=True)
    consulter_cmd.add_argument("--famille", choices=FAMILLES, required=True)
    consulter_cmd.add_argument("--debut", help="AAAA-MM. Avec --fin, renvoie le relevé au lieu de la dernière photo.")
    consulter_cmd.add_argument("--fin", help="AAAA-MM.")
    consulter_cmd.add_argument("--sortie")
    consulter_cmd.add_argument("--s3", action="store_true")
    return analyseur


def main(argv: list[str] | None = None) -> int:
    _stdio_utf8()
    args = construire_analyseur().parse_args(argv)
    try:
        _preparer(args)
        if args.commande == "exporter":
            return _commande_exporter(args)
        return _commande_consulter(args)
    except ErreurExport as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
